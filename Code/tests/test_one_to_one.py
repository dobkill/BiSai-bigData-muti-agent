from itertools import permutations
from types import SimpleNamespace

import pytest

from fusion_agent.agents.alignment import AlignmentAgent
from fusion_agent.agents.event_units import EventUnitAgent
from fusion_agent.agents.extraction import ExtractionAgent, text_times
from fusion_agent.agents.one_to_one import OneToOneAlignmentAgent
from fusion_agent.agents.verification import VerificationAgent
from fusion_agent.assignment import maximum_assignment
from fusion_agent.config import AlignmentSettings, Settings, VerificationSettings
from fusion_agent.evidence_store import EvidenceStore
from fusion_agent.models import DayResult, Document, Observation


def observation(oid, source, platform="054A", **fields):
    facts = {"country": ["中国"], "platform": [platform], "direction": ["东海"], "region": ["东海中部"],
             "start": ["2025-05-01T10:00:00"], "end": ["2025-05-01T12:00:00"], "quantity": [1]}
    facts.update(fields)
    return Observation(id=oid, date="2025-05-01", source=source, path=oid + ".txt", facts=facts,
                       metadata={"platform_kinds": ["ship"]})


def test_assignment_reconsiders_first_choice_for_the_overall_solution():
    assert maximum_assignment([[9, 8], [9, 1]]) == ([1, 0], 17)
    assert maximum_assignment([[0, 0], [0, 5]]) == ([None, 1], 5)
    assert maximum_assignment([[9, 8], [9, 1]], {(0, 1)}) == ([0, 1], 10)


@pytest.mark.parametrize("weights", [
    [[1, 8, 2], [7, 3, 4], [2, 6, 9]],
    [[7, 0, 5], [7, 0, 1], [3, 8, 0], [0, 2, 0]],
    [[0, -1], [4, 4]],
])
def test_solver_matches_exhaustive_small_assignments(weights):
    n, width = len(weights), len(weights[0])
    expected = max(sum(max(0, weights[i][j]) if j < width else 0 for i, j in enumerate(choice))
                   for choice in permutations(range(width + n), n))
    result, objective = maximum_assignment(weights)
    assert objective == expected
    assert len({j for j in result if j is not None}) == sum(j is not None for j in result)


def test_main_groups_are_news_anchored_and_recompute_stable_ids(monkeypatch):
    observations = [observation("a", "open_source"), observation("b", "open_source"),
                    observation("x", "opinion"), observation("y", "opinion")]
    weights = {frozenset(("a", "x")): 9, frozenset(("a", "y")): 8,
               frozenset(("b", "x")): 9, frozenset(("b", "y")): 1}
    monkeypatch.setattr(OneToOneAlignmentAgent, "pair_score", lambda self, a, b: weights[frozenset((a.id, b.id))])
    agent = AlignmentAgent(AlignmentSettings(minimum_pair_score=1))
    first = agent.run(observations)
    assert {c.assignment["news_id"]: c.assignment["opinion_id"] for c in first} == {"a": "y", "b": "x"}
    assert {c.id: c.observation_ids for c in agent.run(list(reversed(observations)))} == {c.id: c.observation_ids for c in first}
    weights[frozenset(("b", "y"))] = 30
    second = agent.run(observations)
    assert {c.id for c in first} == {c.id for c in second}
    assert {c.assignment["news_id"]: c.assignment["opinion_id"] for c in second} == {"a": "x", "b": "y"}


def test_equal_radar_alternatives_remain_in_one_group_and_an_unassigned_list():
    observations = [observation("n", "open_source"), observation("p", "opinion"),
                    observation("r1", "radar"), observation("r2", "radar")]
    agent = AlignmentAgent(AlignmentSettings())
    groups = agent.run(observations)
    assert len(groups) == 1 and groups[0].relation_status == "candidate"
    assert len(groups[0].assignment["radar_candidates"]) == 2
    assert len(agent.unassigned_observation_ids) == 1
    assert "multiple_compatible_hypotheses" in groups[0].warnings


def test_last_remaining_unrelated_material_is_not_forced_into_a_pair():
    observations = [observation("n", "open_source"),
                    observation("p", "opinion", platform="055", country=["日本"], start=[], end=[])]
    agent = AlignmentAgent(AlignmentSettings())
    groups = agent.run(observations)
    assert len(groups) == 1 and groups[0].observation_ids == ["n"]
    assert groups[0].relation_status == "unmatched"
    assert agent.unassigned_observation_ids == ["p"]


@pytest.mark.parametrize("other_source", ["opinion", "radar"])
def test_missing_opinion_or_radar_keeps_link_but_makes_news_event_untrusted(other_source):
    observations = [observation("n", "open_source"), observation("other", other_source)]
    groups = AlignmentAgent(AlignmentSettings()).run(observations)
    assert groups[0].relation_status == "confirmed"
    event = VerificationAgent(VerificationSettings()).run(groups, observations)[0]
    assert event.classification == "untrusted" and not event.accepted
    assert "three_sources_not_linked" in event.verification_reasons


def test_identity_conflict_can_be_paired_but_never_accepted_as_a_fact():
    observations = [observation("n", "open_source", country=["意大利"], start=[], end=[]),
                    observation("p", "opinion", start=[], end=[]), observation("r", "radar")]
    group = AlignmentAgent(AlignmentSettings()).run(observations)[0]
    assert set(group.observation_ids) == {"n", "p", "r"}
    assert group.relation_status == "conflicted"
    event = VerificationAgent(VerificationSettings()).run([group], observations)[0]
    assert not event.accepted and "fact_conflict:country" in event.reasons


def test_nearby_hours_triangle_is_not_confirmed_by_the_two_other_edges():
    observations = [observation("n", "open_source"), observation("p", "opinion"),
                    observation("r", "radar", start=["2025-05-01T14:00:00"], end=["2025-05-01T15:00:00"])]
    group = AlignmentAgent(AlignmentSettings()).run(observations)[0]
    assert group.relation_status == "candidate"
    event = VerificationAgent(VerificationSettings()).run([group], observations)[0]
    assert not event.accepted and event.verdicts["start"].status == "insufficient"


def joint_activity():
    texts = [Document(id="n", date="2025-05-01", source="open_source", path="001.txt",
                      text="正文：2025年5月1日，中国1艘039A型潜艇和1艘056A型护卫舰在东海中部从10时00分至12时00分执行反潜巡逻。"),
             Document(id="p", date="2025-05-01", source="opinion", path="048.txt",
                      text="正文：2025年5月1日，中国1艘039A型潜艇在东海执行反潜巡逻。")]
    extractor = ExtractionAgent(Settings(), None)
    raw = [extractor.text(d) for d in texts]
    raw += [observation("r1", "radar", "039A", task=["反潜巡逻"]),
            observation("r2", "radar", "056A", task=["反潜巡逻"])]
    return texts, EventUnitAgent().run(texts, raw, keep_documents=True)


def test_joint_report_keeps_two_radar_batches_and_verifies_platforms_separately():
    docs, observations = joint_activity()
    assert len(observations) == 4
    groups = AlignmentAgent(AlignmentSettings()).run(observations, docs)
    assert len(groups) == 1 and len(groups[0].observation_ids) == 4
    event = VerificationAgent(VerificationSettings()).run(groups, observations)[0]
    assert event.accepted
    assert set(event.platform_verdicts) == {"039A", "056A"}
    assert set(event.verdicts["platform"].support_sources) == {"open_source", "radar"}
    assert "p" not in event.verdicts["platform"].evidence_ids
    assert "quantity" not in event.facts  # no false formation total of one ship
    assert event.platform_verdicts["056A"]["quantity"].value == 1
    assert all(set(v.evidence_ids) <= {"n", "p", "r1", "r2"}
               for vs in event.platform_verdicts.values() for v in vs.values())


def test_joint_platform_quantity_conflict_cannot_hide_in_other_platform():
    docs, observations = joint_activity()
    observations[-1].facts["quantity"] = [2]
    groups = AlignmentAgent(AlignmentSettings()).run(observations, docs)
    event = VerificationAgent(VerificationSettings()).run(groups, observations)[0]
    assert not event.accepted and event.platform_verdicts["056A"]["quantity"].status == "conflict"
    assert "fact_conflict:quantity" in event.reasons


def test_quoted_llm_review_cannot_turn_equivalent_assignments_into_confirmed_facts():
    class Reviewer:
        settings = SimpleNamespace(fail_on_error=True)
        def complete_json(self, *args):
            return {"opinion_id": "p2", "decision": "same_event", "news_quote": "中国054A", "opinion_quote": "中国054A", "reason": "同平台"}
    observations = [observation("n", "open_source"), observation("p1", "opinion"), observation("p2", "opinion")]
    docs = [Document(id=o.id, date=o.date, source=o.source, path=o.path, text="中国054A在东海活动") for o in observations]
    agent = AlignmentAgent(AlignmentSettings(), Reviewer())
    group = agent.run(observations, docs)[0]
    assert group.assignment["opinion_id"] == "p2"
    assert group.relation_status == "candidate" and group.assignment["evidence_margin"] == 0
    assert agent.audit["llm_validated_reviews"] == 1


def test_unanchored_llm_review_is_ignored():
    class Reviewer:
        settings = SimpleNamespace(fail_on_error=True)
        def complete_json(self, *args):
            return {"opinion_id": "p2", "decision": "same_event", "news_quote": "不存在的证据", "opinion_quote": "不存在的证据", "reason": ""}
    observations = [observation("n", "open_source"), observation("p1", "opinion"), observation("p2", "opinion")]
    docs = [Document(id=o.id, date=o.date, source=o.source, path=o.path, text="中国054A在东海活动") for o in observations]
    agent = AlignmentAgent(AlignmentSettings(), Reviewer())
    group = agent.run(observations, docs)[0]
    assert group.relation_status == "candidate" and agent.audit["llm_unanchored_reviews"] == 1


def test_store_preserves_normalized_places_and_previous_assignment(tmp_path):
    obs = [observation("n", "open_source"), observation("p", "opinion")]
    store = EvidenceStore(tmp_path / "evidence.sqlite")
    try:
        store.replace_day("2025-05-01", obs, [])
        assert store.observations("2025-05-01")[0].facts["direction"] == ["东海"]
        first = DayResult(date="2025-05-01", observations=obs, candidates=[], events=[], trace=[])
        store.record(first)
        second = first.model_copy(deep=True)
        second.unassigned_observation_ids = ["p"]
        store.record(second)
        assert store.latest("2025-05-01").unassigned_observation_ids == ["p"]
        assert store.connection.execute("SELECT COUNT(*) FROM assignment_runs").fetchone()[0] == 2
        assert store.connection.execute("SELECT COUNT(*) FROM assignment_runs WHERE previous_id IS NOT NULL").fetchone()[0] == 1
        store.record(first)
        assert store.latest("2025-05-01").unassigned_observation_ids == []
        assert store.connection.execute("SELECT COUNT(*) FROM assignment_runs").fetchone()[0] == 3
    finally:
        store.close()


def test_separated_activity_boundaries_are_parsed():
    starts, ends = text_times("2025年5月1日，自11时26分起执行任务至13时51分。", "2025-05-01")
    assert starts == ["2025-05-01T11:26:00"] and ends == ["2025-05-01T13:51:00"]


def test_named_other_regions_keep_explicit_conflicts():
    from fusion_agent.domain import REGION_DIRECTION, regions
    assert regions("在鄂霍次克海活动") == ["鄂霍次克海"]
    assert regions("在中太平洋海域活动") == ["中太平洋"]
    assert REGION_DIRECTION["夏威夷附近海空域"] == "其他"


def test_main_export_preserves_unassigned_sources_without_extra_groups(tmp_path):
    from test_pipeline import dataset
    from fusion_agent.config import Paths, ReportSettings
    from fusion_agent.pipeline import AgentPipeline
    import json
    dataset(tmp_path / "data")
    extra = tmp_path / "data" / "2025.5.10" / "网上舆论" / "002.txt"
    extra.write_text("标题：美国F-35B在地中海活动\n正文：美国F-35B在地中海活动。", encoding="utf-8")
    import openpyxl
    workbook_path = tmp_path / "data/2025.5.10/雷达.xlsx"
    workbook = openpyxl.load_workbook(workbook_path)
    for _ in range(2):
        workbook["舰船"].append(["SHIP-2", "日本", "摩耶级驱逐舰", "黄海", "2025.05.10 14:00:00", "2025.05.10 15:00:00", 1, "联合巡逻"])
    workbook.save(workbook_path)
    workbook.close()
    settings = Settings(paths=Paths(data_dir=tmp_path / "data", output_dir=tmp_path / "out", cache_dir=tmp_path / "cache"),
                        report=ReportSettings(lookback_days=1))
    pipeline = AgentPipeline(settings)
    try:
        manifest = pipeline.run(["2025-05-10"])
        assert manifest["total_candidates"] == 1 and manifest["unassigned_observations"] == 2
        index = json.loads((tmp_path / "out/review/2025-05-10/索引.json").read_text(encoding="utf-8"))
        assert len(index["groups"]) == 1 and len(index["unassigned_observations"]) == 2
        assert index["counts"] == {"可信事件": 1, "不可信事件": 0}
        assert index["unpaired_record_count"] == 2 and len(index["unassigned_records"]) == 2
        for record in index["unassigned_records"]:
            folder = tmp_path / "out/review" / record["path"]
            assert record["status"] == "不可信事件" and folder.parent.name == "不可信事件"
            if "opinion" in record["source_types"].values():
                assert (folder / "舆论.txt").read_bytes() == extra.read_bytes()
            else:
                workbook = openpyxl.load_workbook(folder / "雷达.xlsx", read_only=True, data_only=True)
                try:
                    rows = list(workbook["舰船"].values)
                    assert len(rows) == 3 and all(row[1] == "SHIP-2" for row in rows[1:])
                finally:
                    workbook.close()
        assert manifest["review_export"]["unassigned_record_sources"] == {"radar": 1, "opinion": 1}
        assert manifest["review_export"]["unique_observations"]["opinion"] == 2
        # Validate fidelity across both the paired group and standalone copies.
        import importlib.util
        from pathlib import Path
        specification = importlib.util.spec_from_file_location(
            "validate_review", Path(__file__).resolve().parents[1] / "scripts/validate_one_to_one.py")
        validation_module = importlib.util.module_from_spec(specification)
        specification.loader.exec_module(validation_module)
        validation = validation_module.validate_day(settings.paths.output_dir / "daily/2025-05-10.json",
                                                     tmp_path / "out/review", settings.paths.data_dir)
        assert validation["exported_original_texts"] == 3
        assert validation["assigned_radar_rows"] == 2 and validation["exported_radar_rows"] == 4
    finally:
        pipeline.close()
