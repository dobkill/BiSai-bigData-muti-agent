import csv
import json

import openpyxl
import pytest

from fusion_agent.config import AlignmentSettings, Paths, ReportSettings, Settings
from fusion_agent.pipeline import AgentPipeline
from fusion_agent.review_export import ReviewExporter, export_saved_results
from fusion_agent.utils import write_json


@pytest.fixture
def review_data(tmp_path):
    root = tmp_path / "data"
    folder = root / "2025.5.10"
    for name in ("开源军情", "网上舆论"):
        (folder / name).mkdir(parents=True)
        # 同一新闻/舆论与 SHIP-1 对应，两个记录是轨迹点而非两艘船。
        text = "标题：中国1艘054A型护卫舰在南海执行战备巡逻\n正文：2025年5月10日，中国1艘054A型护卫舰在南海从10时00分至12时00分执行战备巡逻。\n"
        (folder / name / "001.txt").write_bytes(text.encode("utf-8"))
    (folder / "开源军情" / "002.txt").write_text("标题：美国F-35B战斗机在地中海活动\n正文：美国F-35B战斗机在地中海活动。", encoding="utf-8")
    wb = openpyxl.Workbook()
    air = wb.active
    air.title = "飞机"
    air.append(["批号", "国家地区", "目标类型", "位置时间", "经度", "纬度", "目标数量", "任务"])
    ship = wb.create_sheet("舰船")
    ship.append(["批号", "国家地区", "机舰名", "活动区域", "首点时间", "结束时间", "数量", "任务", "发现时间"])
    ship.append(["SHIP-1", "中国", "054A型护卫舰", "南海", "2025.05.10 10:00:00", "2025.05.10 12:00:00", 1, "战备巡逻", "2025.05.10 11:00:00"])
    ship.append(["SHIP-1", "中国", "054A型护卫舰", "南海", "2025.05.10 10:00:00", "2025.05.10 12:00:00", 1, "战备巡逻", "2025.05.10 10:00:00"])
    for time in ("14:00:00", "15:00:00"):
        ship.append(["SHIP-2", "日本", "摩耶级驱逐舰", "黄海", "2025.05.10 14:00:00", "2025.05.10 15:00:00", 1, "联合巡逻", "2025.05.10 " + time])
    wb.save(folder / "雷达.xlsx")
    wb.close()
    settings = Settings(paths=Paths(data_dir=root, output_dir=tmp_path/"out", cache_dir=tmp_path/"cache", review_dir=tmp_path/"review"),
                        report=ReportSettings(lookback_days=1), alignment=AlignmentSettings(mode="legacy"))
    pipeline = AgentPipeline(settings)
    try:
        day = pipeline.process_day("2025-05-10")
    finally:
        pipeline.close()
    return settings, day


def group_folder(settings, day, status):
    groups = json.loads((settings.paths.review_dir/day.date/"索引.json").read_text(encoding="utf-8"))["groups"]
    group = next(g for g in groups if g["status"] == status)
    return settings.paths.review_dir / group["path"]


def test_export_copies_original_bytes_filters_rows_and_keeps_all_observations(review_data):
    settings, day = review_data
    before = {p: p.read_bytes() for p in settings.paths.data_dir.rglob("*") if p.is_file()}
    summary = ReviewExporter(settings.paths.data_dir, settings.paths.review_dir).export([day])
    assert summary["counts"] == {"可信事件": 1, "可疑事件": 0, "未配对记录": 2}
    assert summary["unique_observations"] == {"radar": 2, "open_source": 2, "opinion": 1}
    folder = group_folder(settings, day, "可信事件")
    assert (folder/"新闻.txt").read_bytes() == before[settings.paths.data_dir/"2025.5.10"/"开源军情"/"001.txt"]
    assert (folder/"舆论.txt").read_bytes() == before[settings.paths.data_dir/"2025.5.10"/"网上舆论"/"001.txt"]
    wb = openpyxl.load_workbook(folder/"雷达.xlsx", read_only=True, data_only=True)
    try:
        rows = list(wb["舰船"].values)
        assert len(rows) == 3
        assert {r[1] for r in rows[1:]} == {"SHIP-1"}
        assert [r[0] for r in rows[1:]] == [3, 2]  # 按时间排序并保留原始行号
        assert "来源说明" in wb.sheetnames
    finally:
        wb.close()
    assert "新闻 ↔ 舆论" in (folder/"关联说明.md").read_text(encoding="utf-8")
    assert "国家地区" in (folder/"关联说明.md").read_text(encoding="utf-8")
    readme = (folder/"README.md").read_text(encoding="utf-8")
    assert "分组依据" in readme and "活动区间重叠或包含" in readme
    assert "匹配分数" not in readme and len(readme.splitlines()) < 40
    overview = (settings.paths.review_dir/"README.md").read_text(encoding="utf-8")
    assert "缺失记未知" in overview and "仅保留候选" in overview
    assert all(p.read_bytes() == data for p, data in before.items())
    with (settings.paths.review_dir/"索引.csv").open(encoding="utf-8-sig", newline="") as f:
        assert len(list(csv.DictReader(f))) == 3


def test_rerun_preserves_review_and_moves_status_with_version_notice(review_data):
    settings, day = review_data
    exporter = ReviewExporter(settings.paths.data_dir, settings.paths.review_dir)
    exporter.export([day])
    old_folder = group_folder(settings, day, "可信事件")
    review = old_folder/"人工复核.csv"
    edited = review.read_bytes().replace("待复核".encode("utf-8"), "人工通过".encode("utf-8"))
    review.write_bytes(edited)
    exporter.export([day])
    assert review.read_bytes() == edited
    assert not json.loads((old_folder/"事件信息.json").read_text(encoding="utf-8"))["review_record_stale"]
    event = next(e for e in day.events if e.accepted)
    event.accepted = False
    event.reasons = ["fact_conflict:quantity"]
    event.classification = "suspicious"
    event.verification_reasons = event.reasons[:]
    summary = exporter.export([day])
    new_folder = group_folder(settings, day, "可疑事件")
    assert not old_folder.exists()
    assert (new_folder/"人工复核.csv").read_bytes() == edited
    assert summary["counts"]["可疑事件"] == 1
    assert json.loads((new_folder/"事件信息.json").read_text(encoding="utf-8"))["review_record_stale"]
    assert "数量存在显式冲突" in (new_folder/"关联说明.md").read_text(encoding="utf-8")


def test_review_saved_by_spreadsheet_in_local_encoding_is_preserved(review_data):
    settings, day = review_data
    exporter = ReviewExporter(settings.paths.data_dir, settings.paths.review_dir)
    exporter.export([day])
    folder = group_folder(settings, day, "可信事件")
    review = folder/"人工复核.csv"
    data = review.read_text(encoding="utf-8-sig").replace("待复核", "人工通过").encode("gb18030")
    review.write_bytes(data)
    exporter.export([day])
    assert review.read_bytes() == data
    assert not json.loads((folder/"事件信息.json").read_text(encoding="utf-8"))["review_record_stale"]


@pytest.mark.parametrize("changed_source", ["text", "radar"])
def test_changed_original_data_cannot_be_exported_as_verified(review_data, changed_source):
    settings, day = review_data
    if changed_source == "text":
        path = settings.paths.data_dir/"2025.5.10"/"开源军情"/"001.txt"
        path.write_text(path.read_text(encoding="utf-8").replace("1艘", "3艘"), encoding="utf-8")
    else:
        path = settings.paths.data_dir/"2025.5.10"/"雷达.xlsx"
        wb = openpyxl.load_workbook(path)
        wb["舰船"].cell(2, 7, 3)
        wb.save(path)
        wb.close()
    with pytest.raises(ValueError, match="已改变"):
        ReviewExporter(settings.paths.data_dir, settings.paths.review_dir).export([day])
    assert not settings.paths.review_dir.exists()


def test_output_and_source_path_guards(review_data, tmp_path):
    settings, day = review_data
    with pytest.raises(ValueError, match="原始数据目录分开"):
        ReviewExporter(settings.paths.data_dir, settings.paths.data_dir/"review")
    outside = tmp_path/"outside.txt"
    outside.write_text("private", encoding="utf-8")
    day.observations[0].path = "../outside.txt"
    with pytest.raises(ValueError, match="越过数据目录"):
        ReviewExporter(settings.paths.data_dir, settings.paths.review_dir).export([day])


def test_export_saved_results_reads_all_saved_dates_without_pipeline(review_data):
    settings, day = review_data
    daily = settings.paths.output_dir/"daily"
    write_json(daily/(day.date+".json"), day.model_dump(mode="json"))
    result = export_saved_results(settings.paths.data_dir, daily, settings.paths.review_dir)
    assert result["exported_dates"] == [day.date]
    assert result["total_groups"] == len(day.candidates)
    with pytest.raises(FileNotFoundError, match="关联结果缺失"):
        export_saved_results(settings.paths.data_dir, daily, settings.paths.review_dir, ["2025-05-11"])


def test_removed_groups_are_archived_with_manual_files_preserved(review_data):
    settings, day = review_data
    exporter = ReviewExporter(settings.paths.data_dir, settings.paths.review_dir)
    exporter.export([day])
    folder = group_folder(settings, day, "可信事件")
    (folder/"我的笔记.txt").write_text("需要队友再看一下", encoding="utf-8")
    accepted = next(e for e in day.events if e.accepted)
    removed = set(accepted.evidence_ids)
    day.events = [e for e in day.events if e.id != accepted.id]
    day.candidates = [c for c in day.candidates if c.id != accepted.id]
    day.observations = [o for o in day.observations if o.id not in removed]
    exporter.export([day])
    history = settings.paths.review_dir/day.date/"历史分组"/folder.name
    assert not folder.exists()
    assert (history/"我的笔记.txt").read_text(encoding="utf-8") == "需要队友再看一下"


def test_old_category_migrates_without_losing_manual_review_or_notes(review_data):
    settings, day = review_data
    exporter = ReviewExporter(settings.paths.data_dir, settings.paths.review_dir)
    exporter.export([day])
    folder = group_folder(settings, day, "可信事件")
    review = (folder / "人工复核.csv").read_bytes().replace("待复核".encode(), "人工通过".encode())
    (folder / "人工复核.csv").write_bytes(review)
    (folder / "我的笔记.txt").write_text("保留人工笔记", encoding="utf-8")
    old_folder = settings.paths.review_dir / day.date / "可确认关联" / folder.name
    old_folder.parent.mkdir()
    folder.rename(old_folder)
    index_path = settings.paths.review_dir / day.date / "索引.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    moved = next(g for g in index["groups"] if g["id"] == next(e.id for e in day.events if e.accepted))
    moved.update(status="可确认关联", path=old_folder.relative_to(settings.paths.review_dir).as_posix())
    write_json(index_path, index)
    # Simulate a saved snapshot predating event classification.
    for event in day.events:
        event.classification = None
    exporter.export([day])
    assert (folder / "人工复核.csv").read_bytes() == review
    assert (folder / "我的笔记.txt").read_text(encoding="utf-8") == "保留人工笔记"
    assert not old_folder.parent.exists()
    assert {p.name for p in (settings.paths.review_dir / day.date).iterdir() if p.is_dir()} == {"可信事件", "可疑事件", "未配对记录"}
