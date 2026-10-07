import csv
import hashlib
import html
import json
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

from .data import sha256
from .geography import REGIONS
from .ollama_client import VERSION, dump


def csv_write(path, rows, headers):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def radar_workbook(path, activities):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheetname in ["飞机", "舰船"]:
        selected = [a for a in activities if a["sheet"]==sheetname]
        if not selected:
            continue
        sheet = book.create_sheet(sheetname)
        headers = selected[0]["headers"]
        sheet.append(["原始文件", "原始工作表", "原始Excel行号", "候选活动ID", "同批号其他候选", *headers])
        for activity in selected:
            for record in activity["records"]:
                sheet.append([Path(activity["path"]).name, sheetname, record["row"], activity["id"],
                              ";".join(activity["sibling_ids"]), *[record["values"][h] for h in headers]])
        sheet.freeze_panes = "F2"
        sheet.auto_filter.ref = sheet.dimensions
    if not book.sheetnames:
        book.create_sheet("无已关联雷达").append(["本组没有已关联雷达，未用不相关材料补齐。"])
    book.save(path)


def safe_name(title):
    import re
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).rstrip(". ")[:55]


def group_markdown(group, news, opinion, radar):
    lines = [f"# {group['id']}：{news['title']}", "", f"**分类：{group['status']}**", "",
             "这是本地模型与程序约束的初步分组结果，未经过人工真值评估。", "",
             "## 已关联材料", "", f"- 新闻：[{news['filename']}](新闻.txt)",
             f"- 舆论：[{opinion['filename']}](舆论.txt)" if opinion else "- 舆论：缺失",
             f"- 雷达：[雷达.xlsx](雷达.xlsx)，{len(radar)}个候选活动，{group['radar_row_count']}条记录", "",
             "## 原文地点与归一化", "", "| 来源 | 原始地点 | 标准区域 | 活动方向 | 精度 |", "|---|---|---|---|---|"]
    for source in [news]+([opinion] if opinion else [])+radar:
        loc = source["location"]
        lines.append(f"| {source['id']} | {loc.get('raw') or '未说明'} | {loc.get('region') or '未说明'} | {loc.get('direction') or '未确认'} | {loc.get('precision')} |")
    lines += ["", "飞机坐标仅保留原始坐标与辅助召回方向；没有官方空间边界时不当作精确海域真值。", "",
              "## 分类依据", "", group["verification"]["explanation"], ""]
    lines += [f"- {reason}" for reason in group["reasons"]] or ["- 通过三源关联、两类来源印证及明确冲突检查。"]
    lines += ["", "## 事实核验", "", "| 平台索引 | 字段 | 结论 | 状态 | 支持来源 | 反对来源 | 说明 |", "|---|---|---|---|---|---|---|"]
    for fact in group["verification"]["facts"]:
        cells = [str(fact["platform_index"]),fact["field"],fact["claim"],fact["status"],",".join(fact["supporting_sources"]),",".join(fact["conflicting_sources"]),fact["reason"]]
        lines.append("| " + " | ".join(c.replace("|","／").replace("\n"," ") for c in cells) + " |")
    lines += ["", "## 雷达归属与异常", ""]
    for activity in radar:
        lines += [f"- {activity['id']}，批号 `{activity['batch']}`，工作表 {activity['sheet']}，原始行：{', '.join(str(r['row']) for r in activity['records'])}"]
        lines += [f"  - {a}" for a in activity["anomalies"]]
        if activity["sibling_ids"]:
            lines += [f"  - 同批号其他候选：{', '.join(activity['sibling_ids'])}，见运行记录或未匹配材料，未删除。"]
    lines += ["", "## 进一步检查", "", "[关联候选.csv](关联候选.csv) 保留被召回材料及模型判断；[结果.json](结果.json) 包含平台归属和完整核验。", ""]
    return "\n".join(lines)


def audit(groups, texts, radars, radar_manifest, output):
    text_owners, radar_owners, row_owners = Counter(), Counter(), Counter()
    for group in groups:
        text_owners[group["news_id"]]+=1
        if group["opinion_id"]:
            text_owners[group["opinion_id"]]+=1
        for rid in group["radar_ids"]:
            radar_owners[rid]+=1
    unmatched_text = [s for s in texts if not text_owners[s["id"]]]
    unmatched_radar = [r for r in radars if not radar_owners[r["id"]]]
    for source in unmatched_text:
        text_owners[source["id"]]+=1
    for activity in unmatched_radar:
        radar_owners[activity["id"]]+=1
    for activity in radars:
        for record in activity["records"]:
            row_owners[(activity["sheet"],record["row"])]+=radar_owners[activity["id"]]
    errors = []
    if set(text_owners) != {s["id"] for s in texts} or any(v != 1 for v in text_owners.values()):
        errors.append("文本没有完整且唯一归属")
    if set(radar_owners) != {r["id"] for r in radars} or any(v != 1 for v in radar_owners.values()):
        errors.append("雷达候选活动没有完整且唯一归属")
    if len(row_owners)!=radar_manifest["rows"] or any(v != 1 for v in row_owners.values()):
        errors.append("原始雷达行没有完整且唯一归属")
    for source in texts:
        if sha256(Path(source["path"])) != source["sha256"]:
            errors.append(f"原始文本发生变化：{source['id']}")
    if sha256(Path(radar_manifest["path"]))!=radar_manifest["sha256"]:
        errors.append("原始雷达工作簿发生变化")
    # Validate the exported copies, not just in-memory assignment.
    exported_texts = 0
    for group in groups:
        directory = output / group["directory"]
        sources = [(group["news_id"],"新闻.txt")]+([(group["opinion_id"],"舆论.txt")] if group["opinion_id"] else [])
        byid = {s["id"]:s for s in texts}
        for sid,name in sources:
            if sha256(directory/name)!=byid[sid]["sha256"]:
                errors.append(f"导出文本字节改变：{sid}")
            exported_texts+=1
    for source in unmatched_text:
        target = output / "未匹配材料" / "文本" / (source["id"]+"_"+source["filename"])
        if sha256(target)!=source["sha256"]:
            errors.append(f"未匹配文本字节改变：{source['id']}")
        exported_texts+=1
    original_records = {(a["sheet"],r["row"]):r["values"] for a in radars for r in a["records"]}
    exported_rows = Counter()
    for path in list(output.glob("*事件/*/雷达.xlsx")) + [output/"未匹配材料"/"雷达.xlsx"]:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in book:
            if sheet.title not in {"飞机","舰船"}:
                continue
            rows=list(sheet.iter_rows(values_only=True))
            headers=rows[0][5:]
            for row in rows[1:]:
                key=(row[1],row[2])
                exported_rows[key]+=1
                if dict(zip(headers,row[5:])) != original_records.get(key):
                    errors.append(f"导出雷达字段变化：{key}")
        book.close()
    if set(exported_rows)!=set(original_records) or any(v!=1 for v in exported_rows.values()):
        errors.append("导出雷达行不完整或重复")
    result={"passed":not errors,"errors":errors,"original_texts":len(texts),"exported_texts":exported_texts,
            "original_radar_rows":radar_manifest["rows"],"exported_radar_rows":sum(exported_rows.values()),
            "unique_radar_rows":len(exported_rows),"unmatched_texts":len(unmatched_text),"unmatched_radar_activities":len(unmatched_radar),
            "original_hashes_unchanged":not any("发生变化" in e for e in errors)}
    dump(output/"运行记录"/"完整性检查.json",result)
    if errors:
        raise RuntimeError("导出完整性检查失败："+"；".join(errors))
    return result


def export_results(groups, texts, radars, radar_manifest, output, args, vector_manifest, client):
    byid = {s["id"]:s for s in texts+radars}
    existing = output / "manifest.json"
    old_groups = json.loads((output/"全部匹配结果.json").read_text(encoding="utf-8")) if existing.exists() and (output/"全部匹配结果.json").exists() else []
    # Preserve prior exports on reruns rather than deleting user review material.
    if old_groups:
        backup = output/"历史运行"/datetime.now().strftime("%Y%m%d_%H%M%S")
        for old in old_groups:
            source = (output/old["directory"]).resolve()
            if source.exists():
                if output.resolve() not in source.parents:
                    raise ValueError("历史导出路径越出日期目录")
                target=backup/old["directory"]
                target.parent.mkdir(parents=True,exist_ok=True)
                shutil.move(str(source),str(target))
        remaining=(output/"未匹配材料").resolve()
        if remaining.exists():
            if output.resolve() not in remaining.parents:
                raise ValueError("历史未匹配材料路径越出日期目录")
            shutil.move(str(remaining),str(backup/"未匹配材料"))
    summary=[]
    used_text=set()
    used_radar=set()
    for group in groups:
        news=byid[group["news_id"]]
        opinion=byid.get(group["opinion_id"])
        radar=[byid[rid] for rid in group["radar_ids"]]
        directory=output/group["status"]/(group["id"]+"_"+safe_name(news["title"]))
        directory.mkdir(parents=True,exist_ok=True)
        group["directory"]=directory.relative_to(output).as_posix()
        shutil.copyfile(news["path"],directory/"新闻.txt")
        used_text.add(news["id"])
        if opinion:
            shutil.copyfile(opinion["path"],directory/"舆论.txt")
            used_text.add(opinion["id"])
        used_radar.update(group["radar_ids"])
        radar_workbook(directory/"雷达.xlsx",radar)
        dump(directory/"结果.json",group)
        candidate_rows=[]
        for kind in ["opinion","radar"]:
            edges={e["candidate_id"]:e for e in group["alignment"][kind+"_edges"]}
            selected={group["opinion_id"]} if kind=="opinion" else set(group["radar_ids"])
            records={r["id"]:r for r in group["candidates"][kind]}
            for sid,edge in edges.items():
                records.setdefault(sid,{"id":sid})
            for sid,record in records.items():
                edge=edges.get(sid,{})
                candidate_rows.append({"来源类型":kind,"材料ID":sid,"原始文件":Path(byid[sid]["path"]).name,
                    "已选":sid in selected,"召回分数":record.get("score"),"向量余弦":record.get("cosine"),
                    "平台相似":record.get("platform_similarity"),"时间重叠":record.get("time_overlap"),
                    "地点关系":record.get("location_relation"),"模型关联":edge.get("relation","未进入模型相关候选"),
                    "置信度":edge.get("confidence"),"关联平台索引":str(edge.get("platform_indices",[])),"关联依据":edge.get("reason","")})
        csv_write(directory/"关联候选.csv",candidate_rows,["来源类型","材料ID","原始文件","已选","召回分数","向量余弦","平台相似","时间重叠","地点关系","模型关联","置信度","关联平台索引","关联依据"])
        (directory/"README.md").write_text(group_markdown(group,news,opinion,radar),encoding="utf-8")
        fact_rows="".join("<tr>"+"".join("<td>"+html.escape(str(value))+"</td>" for value in
            [f["platform_index"],f["field"],f["claim"],f["status"],", ".join(f["supporting_sources"]),", ".join(f["conflicting_sources"]),f["reason"]])+"</tr>" for f in group["verification"]["facts"])
        reasons="".join("<li>"+html.escape(reason)+"</li>" for reason in group["reasons"])
        opinion_link='<a href="舆论.txt">舆论原文</a>' if opinion else "舆论缺失"
        review_html='''<!doctype html><html lang="zh"><meta charset="utf-8"><title>事件证据</title><style>body{font:16px system-ui;margin:30px;line-height:1.6}table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccd4df;padding:8px;text-align:left}pre{white-space:pre-wrap;background:#f3f6fa;padding:16px}a{color:#155baa}</style>'''
        review_html+=f'<p><a href="../../浏览.html">返回总表</a></p><h1>{html.escape(group["id"]+"："+news["title"])}</h1><p><strong>{group["status"]}</strong> · 雷达 {group["radar_row_count"]} 条</p><p><a href="新闻.txt">新闻原文</a> · {opinion_link} · <a href="雷达.xlsx">雷达记录</a> · <a href="关联候选.csv">关联候选</a> · <a href="结果.json">完整结果</a></p>'
        review_html+=f'<h2>分类依据</h2><p>{html.escape(group["verification"]["explanation"])}</p><ul>{reasons}</ul><h2>事实核验</h2><table><thead><tr><th>平台索引</th><th>字段</th><th>事实</th><th>状态</th><th>支持来源</th><th>反对来源</th><th>说明</th></tr></thead><tbody>{fact_rows}</tbody></table>'
        review_html+=f'<h2>新闻原文</h2><pre>{html.escape(news["text"])}</pre><h2>舆论原文</h2><pre>{html.escape(opinion["text"] if opinion else "缺失")}</pre><p>模型初步结果，未经过人工真值评估。异常、剩余材料和原始引用均保留。</p></html>'
        (directory/"说明.html").write_text(review_html,encoding="utf-8")
        summary.append({"组ID":group["id"],"分类":group["status"],"新闻":news["filename"],"舆论":opinion["filename"] if opinion else "缺失",
                        "新闻ID":news["id"],"舆论ID":group["opinion_id"] or "", "雷达活动ID":";".join(group["radar_ids"]),
                        "雷达记录数":group["radar_row_count"],"活动方向":group["direction"],"可入日报":group["report_eligible"],
                        "原因":"；".join(group["reasons"]),"目录":group["directory"]})
    unmatched_text=[s for s in texts if s["id"] not in used_text]
    unmatched_radar=[r for r in radars if r["id"] not in used_radar]
    remaining=output/"未匹配材料"
    (remaining/"文本").mkdir(parents=True,exist_ok=True)
    for source in unmatched_text:
        shutil.copyfile(source["path"],remaining/"文本"/(source["id"]+"_"+source["filename"]))
    radar_workbook(remaining/"雷达.xlsx",unmatched_radar)
    dump(remaining/"材料清单.json",{"texts":unmatched_text,"radars":unmatched_radar,"reason":"未通过全局唯一关联，保留原始材料，不当作已证明虚假"})
    csv_write(output/"匹配总表.csv",summary,list(summary[0]))
    book=openpyxl.Workbook()
    sheet=book.active;sheet.title="匹配总表"
    sheet.append(list(summary[0]))
    for row in summary:
        sheet.append([row[k] for k in summary[0]])
    sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
    book.save(output/"匹配总表.xlsx")
    dump(output/"全部匹配结果.json",groups)
    integrity=audit(groups,texts,radars,radar_manifest,output)
    counts=Counter(g["status"] for g in groups)
    manifest={"version":VERSION,"input":str(args.input.resolve()),"output":str(output),"model":args.model,
              "embedding_model":args.embedding_model,"groups":len(groups),"trusted":counts["可信事件"],"untrusted":counts["不可信事件"],
              "matched_opinions":len(used_text)-len(groups),"matched_radar_activities":len(used_radar),
              "matched_radar_rows":sum(g["radar_row_count"] for g in groups),"report_eligible":sum(g["report_eligible"] for g in groups),
              "http_calls_this_run":client.calls,"cache_hits_this_run":client.hits,"vector_index":vector_manifest,
              "integrity":integrity,"accuracy":"未进行人工真值评估，分类为模型初步结果",
              "spatial_limit":"飞机坐标方向仅辅助召回；缺乏官方边界时不作为精确地点的独立核验证据"}
    dump(output/"manifest.json",manifest)
    dump(output/"运行记录"/"地点归一化规则.json",{"source":"用户提供赛题PDF第2页","directions":REGIONS,
              "rules":["同义归一化","方向归属不等于具体地点等价","缺失不同于其他","通行区域本身为其他，进入/离开核心海域单独记录"]})
    intro=[f"# {args.input.name} 三源匹配结果", "",f"共 {len(groups)} 个新闻主组：可信 {counts['可信事件']}，不可信 {counts['不可信事件']}。",
           f"已关联舆论 {manifest['matched_opinions']} 篇，雷达候选活动 {len(used_radar)} 个，共 {manifest['matched_radar_rows']} 条记录。",
           "", "[打开可筛选浏览页](浏览.html) · [Excel总表](匹配总表.xlsx) · [CSV总表](匹配总表.csv)", "",
           "每组目录包含新闻原文、舆论原文、雷达原始字段和行号、关联候选、事实核验与分类依据。未匹配材料另行保存。",
           f"原始{integrity['original_texts']}篇文本和{integrity['original_radar_rows']}条雷达行已核对完整、唯一保留；源文件SHA-256和导出字段检查通过。", "",
           "使用本地Qwen 9B进行源内抽取、候选关联与三源核验，BGE-M3生成检索向量，程序执行全局唯一分配及冲突约束。",
           "分类尚未经过人工真值评估；不可信包括缺源、关联歧义、未消解异常和事实冲突，不等于全部已证明虚假。",
           "飞机经纬度的宽泛方向判断仅用于召回，不冒充官方区域边界；未支持的精确地点不能被填入结论。", "",
           "## 分组入口", "", "| 组ID | 分类 | 新闻标题 | 舆论ID | 雷达行数 |", "|---|---|---|---|---|"]
    for g in groups:
        intro.append(f"| [{g['id']}]({g['directory']}/README.md) | {g['status']} | {byid[g['news_id']]['title']} | {g['opinion_id'] or '缺失'} | {g['radar_row_count']} |")
    (output/"README.md").write_text("\n".join(intro)+"\n",encoding="utf-8")
    rows="\n".join(f'<tr data-status="{html.escape(r["分类"])}"><td>{r["组ID"]}</td><td>{r["分类"]}</td><td><a href="{html.escape(r["目录"])}/说明.html">{html.escape(r["新闻"])}</a></td><td>{html.escape(r["舆论"])}</td><td>{r["雷达记录数"]}</td><td>{html.escape(r["原因"])}</td></tr>' for r in summary)
    browser='''<!doctype html><html lang="zh"><meta charset="utf-8"><title>三源匹配结果</title><style>body{font:16px system-ui;margin:30px;background:#f6f8fa;color:#182331}table{border-collapse:collapse;background:white;width:100%}td,th{border:1px solid #ccd4df;padding:10px;text-align:left}input,select{font:inherit;padding:8px;margin:8px}a{color:#155baa}</style><h1>三源匹配结果</h1><p>每篇新闻一个主组；点击新闻进入证据目录。模型初步结果，未经过人工真值评估。</p><select id="status"><option value="">全部分类</option><option>可信事件</option><option>不可信事件</option></select><input id="query" placeholder="搜索平台、地点或原因"><p><a href="匹配总表.xlsx">Excel总表</a> · <a href="README.md">阅读说明</a></p><table><thead><tr><th>ID</th><th>分类</th><th>新闻与分组入口</th><th>舆论</th><th>雷达行数</th><th>原因</th></tr></thead><tbody>'''+rows+'''</tbody></table><script>function filter(){const q=document.querySelector('#query').value.toLowerCase(),s=document.querySelector('#status').value;document.querySelectorAll('tbody tr').forEach(r=>r.hidden=!!((s&&r.dataset.status!==s)||!r.textContent.toLowerCase().includes(q)))}document.querySelector('#query').oninput=filter;document.querySelector('#status').onchange=filter;</script></html>'''
    (output/"浏览.html").write_text(browser,encoding="utf-8")
    return manifest
