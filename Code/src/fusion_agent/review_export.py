"""将已保存的关联结果导出为可人工逐事件阅读的三源资料包。"""
from collections import Counter
import csv
from datetime import date
import hashlib
from io import StringIO
import json
import logging
from pathlib import Path
from os.path import relpath
import re
import tempfile
from typing import Iterable
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Candidate, DayResult, Event, FACT_FIELDS, Observation
from .agents.verification import METHOD as VERIFICATION_METHOD, VerificationAgent
from .config import AlignmentSettings, VerificationSettings
from .relations import RELATION_NAMES
from .utils import stable_id, unique, write_json, write_text

logger = logging.getLogger(__name__)
FORMAT = "fusion_review_v1"
LEGACY_STATUSES = ("已核验", "待人工确认", "未匹配")
ASSOCIATION_STATUSES = ("可确认关联", "相关候选", "疑似相关但有冲突", "未找到关联")
EVENT_NAMES = {"trusted": "可信事件", "untrusted": "不可信事件"}
EVENT_STATUSES = tuple(EVENT_NAMES.values())
OLD_EVENT_STATUSES = ("可疑事件", "未配对记录")
STATUSES = EVENT_STATUSES + OLD_EVENT_STATUSES + ASSOCIATION_STATUSES + LEGACY_STATUSES
SOURCE_NAMES = {"open_source": "新闻", "opinion": "舆论", "radar": "雷达"}
FIELD_NAMES = {"country": "国家地区", "platform": "平台", "direction": "方向", "region": "区域",
               "quantity": "数量", "task": "任务", "start": "开始时间", "end": "结束时间"}
VERDICT_NAMES = {"supported": "获得跨源支持", "conflict": "存在显式冲突",
                 "insufficient": "未知（来源支持不足）", "absent": "未知（未给出）"}
REASON_NAMES = {
    "three_sources_not_linked": "尚未关联齐新闻、舆论、雷达三类来源",
    "three_source_association_incomplete": "三类来源尚未完成同一事件关联",
    "fewer_than_two_source_types": "暂时只有单类来源，无法交叉印证；未配对不代表虚假",
    "unlinked_observation": "该观测尚未关联到其他来源",
    "inconsistent_text_alignment": "新闻与舆论之间的关联评分未达到阈值",
    "observation_reused_across_candidates": "同一材料被关联到多个候选分组，需要人工确认归属",
    "outside_core_or_unknown_direction": "活动在核心海域以外，或核心方向尚未核验，不纳入日报",
    "multiple_platforms_unresolved": "材料包含多个平台，尚未完成子事件拆分",
    "platform_kind_conflict": "材料中的飞机/舰船类别不一致",
    "invalid_time_interval": "开始与结束时间存在异常",
    "radar_date_mismatch": "雷达时间与日期目录不一致",
    "text_date_mismatch": "文本事件日期与日期目录不一致",
    "insufficient_radar_points": "雷达记录点数量不足",
    "radar_cross_direction_track": "雷达轨迹跨越多个方向，需要核验",
    "radar_implausible_track_jump": "雷达轨迹出现不合理的位置跳变",
    "radar_individual_identity_conflict": "已关联的同一目标存在编号或名称冲突，需要核验",
    "multiple_compatible_hypotheses": "存在多个合理的对应组合，已全部保留，需人工确认归属",
    "semantic_opinion_association_uncertain": "模型尚不能确认新闻与舆论属于同一次活动",
    "semantic_radar_association_uncertain": "模型尚不能确认所选雷达属于该活动",
    "semantic_verification_failed": "模型核验失败或证据校验不通过，保持不可信",
}
INDEX_COLUMNS = ["日期", "状态", "事件ID", "三源齐全", "日报接纳", "国家地区", "平台", "方向",
                 "新闻标题", "舆论标题", "雷达批号", "雷达记录数", "需要确认的问题", "分组目录", "分组依据", "字段核验",
                 "新闻编号", "舆论编号", "主组口径"]


def explain(reason: str) -> str:
    prefix, _, value = reason.partition(":")
    if prefix == "fact_conflict":
        return f"{FIELD_NAMES.get(value, value)}存在显式冲突"
    if prefix == "fact_not_verified":
        return f"{FIELD_NAMES.get(value, value)}未获得足够跨源支持"
    if prefix == "missing_source":
        return f"缺少可关联的{SOURCE_NAMES.get(value, value)}材料；未满足三源关联条件，归入不可信"
    if prefix == "ambiguous_alignment":
        return f"{SOURCE_NAMES.get(value, value)}存在评分接近的多个候选，关联有歧义"
    if prefix == "uncertain_coordinate_points":
        return f"{value} 个坐标点位于方向判定边界或未知区域"
    if prefix == "llm_fallback":
        return "模型抽取回退到规则抽取：" + value
    if prefix == "association_not_confirmed":
        return "事件关联尚未确认：" + RELATION_NAMES.get(value, value)
    if prefix == "semantic_association_not_confirmed":
        return "模型未确认同一活动：" + {"different_activity": "材料描述不同活动", "uncertain": "归属证据不足"}.get(value, value)
    if prefix == "semantic_extraction_invalid":
        return "模型抽取未通过证据校验，需重新核验原文"
    if prefix == "llm_unanchored_claim":
        return "模型的" + FIELD_NAMES.get(value, value) + "声明缺少连续原文引用，未采用"
    if prefix == "llm_unvalidated_claim":
        return "模型的" + FIELD_NAMES.get(value, value) + "声明未通过引用校验，未采用"
    return REASON_NAMES.get(reason, reason)


def display(values) -> str:
    if values is None or values == [] or values == "":
        return "未给出"
    if not isinstance(values, list):
        values = [values]
    return "；".join(str(v).removeprefix("generic:") for v in values)


def table_cell(value) -> str:
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def csv_text(columns: list[str], rows: list[list]) -> str:
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(columns)
    for row in rows:
        # Excel 阅读时将潜在公式开头按文本显示，原文文件仍保留全部原始字节。
        writer.writerow(["'" + str(v) if str(v).startswith(("=", "+", "-", "@")) else v for v in row])
    return "\ufeff" + stream.getvalue()


def write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as f:
        f.write(data)
        temp = Path(f.name)
    temp.replace(path)


def json_value(value):
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))


def review_version(path: Path) -> str | None:
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            # Excel/WPS 另存为普通 CSV 时可能使用本地中文编码，保留其原始字节。
            text = data.decode("gb18030")
    return next(csv.DictReader(StringIO(text)), {}).get("数据版本")


class ReviewExporter:
    """只消费 DayResult 和原始文件，不改变匹配算法或模型调用。"""

    def __init__(self, data_dir: Path, output_dir: Path, analysis_dir: Path | None = None):
        self.data_dir = data_dir.resolve()
        self.output_dir = output_dir.resolve()
        self.analysis_dir = analysis_dir.resolve() if analysis_dir else None
        if (self.output_dir.is_relative_to(self.data_dir) or self.data_dir.is_relative_to(self.output_dir)):
            raise ValueError("人工阅读导出目录必须与原始数据目录分开，不能互相包含")
        if self.analysis_dir and (self.output_dir.is_relative_to(self.analysis_dir) or self.analysis_dir.is_relative_to(self.output_dir)):
            raise ValueError("人工阅读导出目录不能覆盖关联结果目录")
        manifest = self.output_dir / "manifest.json"
        if manifest.exists():
            previous = json.loads(manifest.read_text(encoding="utf-8"))
            if previous.get("format") != FORMAT or previous.get("data_dir") != str(self.data_dir):
                raise ValueError("导出目录已属于其他内容或数据集，请另选 review_dir")

    def target(self, path: Path) -> Path:
        resolved = path.resolve()
        if not resolved.is_relative_to(self.output_dir):
            raise ValueError(f"导出路径越过人工阅读目录：{path}")
        return resolved

    def source(self, relative: str) -> Path:
        path = (self.data_dir / relative).resolve()
        if not path.is_relative_to(self.data_dir):
            raise ValueError(f"原始文件路径越过数据目录：{relative}")
        if not path.is_file():
            raise FileNotFoundError(f"找不到原始材料：{path}")
        return path

    def preflight(self, day: DayResult) -> tuple[dict[str, bytes], dict[str, dict]]:
        """先确认原始文件仍对应分析快照，再为一整天的分组准备材料。"""
        date.fromisoformat(day.date)
        observations = {o.id: o for o in day.observations}
        events = {e.id: e for e in day.events}
        if len(observations) != len(day.observations) or len(events) != len(day.events):
            raise ValueError(f"{day.date} 关联结果包含重复 ID")
        ids = [c.id for c in day.candidates]
        if len(set(ids)) != len(ids) or set(ids) != set(events):
            raise ValueError(f"{day.date} 候选和核验事件不一一对应")
        referenced = set()
        for candidate in day.candidates:
            if not candidate.observation_ids or not set(candidate.observation_ids) <= observations.keys():
                raise ValueError(f"{day.date} 候选引用了缺失观测")
            if set(events[candidate.id].evidence_ids) != set(candidate.observation_ids):
                raise ValueError(f"{day.date} 核验证据与候选分组不一致")
            referenced.update(candidate.observation_ids)
        unassigned = set(day.unassigned_observation_ids)
        if len(unassigned) != len(day.unassigned_observation_ids) or referenced & unassigned or not unassigned <= observations.keys():
            raise ValueError("未归属清单重复、与主组重叠或引用缺失观测")
        if referenced | unassigned != set(observations):
            raise ValueError(f"{day.date} 存在未保存到任何候选的材料，请重新运行关联")
        if any(c.assignment for c in day.candidates):
            appearances = Counter(i for c in day.candidates for i in c.observation_ids)
            if any(count != 1 for count in appearances.values()):
                raise ValueError("一一对应主组重复占用了来源材料")
            for candidate in day.candidates:
                sources = Counter(observations[i].source for i in candidate.observation_ids)
                if sources["open_source"] != 1 or sources["opinion"] > 1:
                    raise ValueError("每个主组必须固定一篇新闻，最多配一篇舆论")
            if any(observations[i].source == "open_source" for i in unassigned):
                raise ValueError("新闻必须各自建立主组")
        texts = {}
        workbooks = {}
        radar = {}
        for obs in day.observations:
            if obs.date != day.date:
                raise ValueError("分组中存在跨日期观测")
            path = self.source(obs.path)
            if obs.source != "radar":
                data = path.read_bytes()
                # 分析阶段 read_text 使用通用换行；副本仍保留原始 CRLF/LF 字节。
                text = data.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
                title = re.search(r"^标题[：:]\s*(.*)$", text, re.MULTILINE)
                stored_title = obs.metadata.get("title")
                quotes = [q for qs in obs.quotes.values() for q in qs]
                if ((stored_title and (not title or title.group(1) != stored_title)) or any(q not in text for q in quotes)):
                    raise ValueError(f"原始文本已改变，与关联结果不一致：{obs.path}；请先重新 run")
                texts[obs.id] = data
                continue
            if obs.path not in workbooks:
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    workbooks[obs.path] = {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in wb}
                finally:
                    wb.close()
            sheet = obs.metadata.get("sheet")
            rows = workbooks[obs.path].get(sheet)
            numbers = obs.metadata.get("row_numbers", [])
            raw = obs.metadata.get("raw_rows", [])
            if not rows or not numbers or len(numbers) != len(raw) or len(numbers) != len(set(numbers)):
                raise ValueError(f"雷达观测缺少可追溯的原始记录：{obs.id}")
            headers = list(rows[0])
            selected = []
            for number, snapshot in zip(numbers, raw):
                if not isinstance(number, int) or number < 2 or number > len(rows):
                    raise ValueError(f"雷达行号越界：{obs.path}/{sheet}/{number}")
                values = list(rows[number-1])
                record = dict(zip(headers, values))
                if json_value(record) != json_value(snapshot) or str(record.get("批号")) != str(obs.metadata.get("batch")):
                    raise ValueError(f"原始雷达记录已改变，与关联结果不一致：{obs.path}/{sheet}/{number}；请先重新 run")
                selected.append((number, values))
            # 按观测时间排序，同时保留原始 Excel 行号，方便阅读轨迹。
            time_column = "位置时间" if sheet == "飞机" else "发现时间"
            time_index = headers.index(time_column) if time_column in headers else None
            if time_index is not None:
                selected.sort(key=lambda item: (str(item[1][time_index] or ""), item[0]))
            radar[obs.id] = {"sheet": sheet, "headers": headers, "rows": selected}
        return texts, radar

    @staticmethod
    def status(candidate: Candidate, event: Event, members: list[Observation]) -> str:
        if {o.source for o in members} != {"open_source", "opinion", "radar"}:
            return "不可信事件"
        if event.classification is not None:
            return EVENT_NAMES[event.classification]
        return "可信事件" if event.accepted else "不可信事件"

    @staticmethod
    def fact_status(event: Event) -> str:
        if any(v.status == "conflict" for v in event.verdicts.values()):
            return "存在字段冲突"
        if all(event.verdicts.get(f) and event.verdicts[f].status == "supported" for f in ("country", "platform", "direction")):
            return "关键字段已支持"
        return "关键字段支持不足"

    def group_readme(self, candidate: Candidate, event: Event, members: list[Observation], status: str) -> str:
        source = {o.id: SOURCE_NAMES[o.source] for o in members}
        missing = [SOURCE_NAMES[s] for s in SOURCE_NAMES if s not in {o.source for o in members}]
        lines = [f"# {candidate.date} 分组依据", "", f"事件分类：**{status}**；"
                 f"三源：{'齐全' if not missing else '缺少' + '、'.join(missing)}；"
                 f"字段核验：{self.fact_status(event)}；日报接纳：{'是' if event.accepted else '否'}。", "",
                 "| 对应关系 | 依据 |", "| --- | --- |"]
        for relation in candidate.relations[:6]:
            lines.append(f"| {source[relation.left_id]} ↔ {source[relation.right_id]} | "
                         + table_cell("；".join(relation.basis)) + " |")
        if not candidate.relations:
            lines.append("| 尚无跨源对应 | 未找到足够的身份及时空线索 |")
        pair_basis = {line for r in candidate.relations for line in r.basis}
        overall = [line for line in candidate.basis if line not in pair_basis]
        if overall:
            lines += ["", "组级依据：" + "；".join(overall) + "。"]
        if candidate.assignment:
            a = candidate.assignment
            lines += ["", f"主组：新闻 {a['news_number']} ↔ 舆论 {a.get('opinion_number') or '待配对'}；"
                      f"雷达 {sum(o.source == 'radar' for o in members)} 批。",
                      "备选对应见 [对应候选.csv](对应候选.csv)，备选不另计主组。"]
        # 优先说明具体事实差异，避免缺源和关联状态提示挤掉冲突原因。
        reasons = [r for r in event.reasons if r.startswith("fact_conflict:")]
        reasons += candidate.warnings + [r for r in event.reasons if r != "three_sources_not_linked"]
        notes = unique([explain(w) for w in reasons])
        if notes:
            lines += ["", "需确认：" + "；".join(notes[:4]) + "。"]
        if missing:
            lines += ["", "来源缺失记为未知；未满足三类来源完成同一事件关联的条件，归入不可信。"]
        units = [o for o in members if o.metadata.get("unit_platform")]
        for obs in units:
            lines += ["", f"本组只对应原文中的 **{obs.metadata['unit_platform']}** 子事件；完整原文仍保留。"]
            if obs.metadata.get("unassigned_fields"):
                lines.append("原文中无法明确分配给该平台的字段：" + "、".join(FIELD_NAMES[f] for f in obs.metadata["unassigned_fields"]) + "，未自动补全。")
        for obs in members:
            if obs.metadata.get("parent_radar_id"):
                lines += ["", f"雷达批号 `{obs.metadata['batch']}` 已区分为 {obs.metadata['batch_activity_count']} 个活动观测；"
                          f"本组对应原始行 {obs.metadata['row_numbers']}。其他行分别关联或保留在未归属材料中，不作为本事件的反证。"]
        lines += ["", "打开 [关联说明.md](关联说明.md) 查看三源字段与原始位置；人工意见填写 `人工复核.csv`。", ""]
        return "\n".join(lines)

    def write_radar(self, path: Path, members: list[Observation], records: dict[str, dict]):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        provenance = []
        try:
            for obs in members:
                if obs.source != "radar":
                    continue
                info = records[obs.id]
                name = info["sheet"]
                if name not in wb.sheetnames:
                    ws = wb.create_sheet(name)
                    ws.append(["原始Excel行号"] + info["headers"])
                else:
                    ws = wb[name]
                    if [c.value for c in ws[1]][1:] != info["headers"]:
                        raise ValueError("同名雷达工作表字段不一致")
                for number, values in info["rows"]:
                    ws.append([number] + values)
                provenance.append([obs.id, obs.path, name, obs.metadata["batch"], len(info["rows"]),
                                   ", ".join(str(n) for n in obs.metadata["row_numbers"])])
            source_sheet = wb.create_sheet("来源说明")
            source_sheet.append(["证据ID", "原始工作簿", "原始工作表", "筛选批号", "记录数", "原始Excel行号"])
            for row in provenance:
                source_sheet.append(row)
            for ws in wb:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for row in ws:
                    for cell in row:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                        cell.alignment = Alignment(vertical="top")
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="305496")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[1].height = 30
                for column in ws.columns:
                    length = max((sum(2 if ord(c)>127 else 1 for c in str(cell.value or "")) for cell in column), default=10)
                    ws.column_dimensions[column[0].column_letter].width = min(56, max(14, length+2))
            with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=path.parent, delete=False) as f:
                temp = Path(f.name)
            try:
                wb.save(temp)
                temp.replace(path)
            finally:
                if temp.exists():
                    temp.unlink()
        finally:
            wb.close()

    def description(self, candidate: Candidate, event: Event, members: list[Observation], status: str,
                    files: list[dict], issues: list[str], stale: bool) -> str:
        lines = [f"# 三源分组 {candidate.id}", "", f"日期：{candidate.date}　分类：**{status}**", "",
                 "先关联成事件，再逐字段核验。可信事件要求新闻、舆论、雷达完成同一事件关联，"
                 "关键事实至少两类来源相互支持，第三类来源未给出明确冲突值；其余均归入不可信事件。"
                 "雷达作为时空、数量和轨迹核验依据之一，不作为无条件优先真值。"
                 "核心海域范围决定是否进入日报，单独记录。", "",
                 "## 阅读顺序", "", "先读新闻和舆论原文，再对照雷达记录；最后查看下方字段对照与关联问题。", ""]
        for file in files:
            lines += [f"- [{file['source_name']}：{file['file']}](<{quote(file['file'])}>)。原始位置：`{file['original_path']}`。"]
        if any(o.source == "radar" for o in members):
            lines += ["- [雷达.xlsx](雷达.xlsx)：包含该分组对应雷达活动的原始记录，增加原始 Excel 行号并按观测时间排序。"]
        lines += ["", "## 关联与核验问题", ""]
        if issues:
            lines += ["- " + issue for issue in issues]
        else:
            lines += ["本次核验未发现关联歧义或关键事实冲突。"]
        diagnostic = candidate.assignment.get("opinion_diagnostics", {})
        if diagnostic:
            outcomes = {"selected": "已选定对应舆论", "no_materials": "当天没有舆论材料",
                        "no_candidates_recalled": "代码未召回舆论候选，尚未交给模型审查",
                        "llm_rejected_all": "模型已审查候选舆论，全部判为不同活动",
                        "review_failed_or_not_run": "候选已召回，但模型审查未完成",
                        "global_assignment_not_selected": "存在候选，但整体分配未选定对应"}
            lines += ["", "## 舆论关联过程", "",
                      f"当天舆论 {diagnostic['total_opinions']} 篇；召回 {len(diagnostic['recalled_opinion_ids'])} 篇；模型审查 {len(diagnostic['reviewed_opinion_ids'])} 篇。",
                      outcomes.get(diagnostic["outcome"], diagnostic["outcome"]) + "。",
                      "过滤阶段与具体原因见事件信息.json 的 assignment.opinion_diagnostics。"]
        workflow = candidate.assignment.get("workflow", {})
        if workflow:
            lines += ["", "## 工作流与异常复查", "",
                      "实体解析、多路候选检索、活动关联、整体分配和事实核验分阶段执行；复查提议不直接改变可信度。"]
            for step in workflow.get("recovery", []):
                lines += [f"- 第 {step['round']} 轮复查{SOURCE_NAMES[step['source']]}：读取 {len(step['inspected_ids'])} 个原始候选，提议 {len(step['proposed_ids'])} 个候选重新审查。"]
            axes = event.semantic_review.get("workflow_diagnostics", {})
            if axes:
                names = {"complete": "三源齐全", "missing_source": "尚未关联齐三源", "same_activity": "同一次活动",
                         "different_activity": "不同活动", "uncertain": "未确认", "conflict": "存在冲突",
                         "supported": "关键字段已支持", "insufficient": "支持不足"}
                lines += ["- 来源覆盖：" + names[axes["source_coverage"]] + "；活动关联：" + names[axes["association"]]
                          + "；事实核验：" + names[axes["facts"]] + "。"]
                for source, coverage in axes["search_coverage"].items():
                    if not any(o.source == source for o in members):
                        lines += [f"- {SOURCE_NAMES[source]}已检查 {len(coverage['inspected_ids'])}/{coverage['corpus_size']} 份材料；"
                                  + ("该日候选检查已完成。" if coverage["complete"] else "搜索尚未覆盖全部材料，不能据此确认原始材料不存在。")]
        normalizations = [(SOURCE_NAMES[o.source], entry) for o in members for entry in o.metadata.get("entity_normalizations", [])]
        if normalizations:
            lines += ["", "## 实体名称规范化", ""]
            for source, entry in normalizations:
                lines += [f"- {source}：模型值“{entry['model_value']}” → 标准值“{entry['canonical_value']}”；原始值与对应引文保留在事件信息.json。"]
        type_changes = [(SOURCE_NAMES[o.source], entry) for o in members for entry in o.metadata.get("entity_type_resolutions", [])]
        if type_changes:
            lines += ["", "## 平台类型复核", ""]
            for source, entry in type_changes:
                lines += [f"- {source}：{entry['model_value']} → {entry['canonical_value']}；依据原文：{entry['quote']}。"]
        rechecks = [(SOURCE_NAMES[o.source], check) for o in members for check in o.metadata.get("source_rechecks", [])]
        if rechecks:
            lines += ["", "## 源内字段完整性复核", "", "只读取各来源自己的原文，不借用其他来源补事实。原始抽取及复核回复均保留。"]
            for source, check in rechecks:
                for unit, fields in check["response"].items():
                    for field, decision in fields.items():
                        if decision["claims"]:
                            lines += [f"- {source}活动 {unit}，{FIELD_NAMES[field]}：{claim['value']}；原文：{claim['quote']}。"
                                      for claim in decision["claims"]]
                        elif decision["status"] == "uncertain":
                            lines += [f"- {source}活动 {unit}，{FIELD_NAMES[field]}：角色未确认，保持未知。"]
        if event.semantic_review:
            review = event.semantic_review
            engine = "模型语义核验及代码证据检查" if review.get("method") != "offline_rules" else "离线规则基线"
            lines += ["", "## 判断来源", "", f"本事件判断方式：{engine}。"]
            if review.get("response"):
                response = review["response"]
                association_label = {"same_activity": "同一次活动", "different_activity": "不同活动", "uncertain": "尚不确定"}[response["association"]]
                lines += ["以下保留模型原始说明，便于与原文对照；最终字段状态、事实值和支持来源见下方核验结论。",
                          f"- 模型活动归属说明：{association_label}；{response['reason']}"]
                lines += [f"- 模型{FIELD_NAMES[f['field']]}说明：{f['reason']}" for f in response["fields"]]
                if response.get("excluded_claims"):
                    lines += [f"- 排除候选 {e['claim_id']}：{e['role']}；{e['reason']}" for e in response["excluded_claims"]]
            if review.get("status") == "failed":
                lines += ["- 模型核验失败，保持不可信；详见事件信息中的失败原因。"]
        if candidate.relation_status is not None:
            source = {o.id: SOURCE_NAMES[o.source] for o in members}
            lines += ["", "## 关联判定依据", "", "强字段冲突不能靠其他信息抵消；缺失记为未知。关联状态与字段核验分别记录。", "",
                      "| 来源关系 | 判定 | 具体依据 |", "| --- | --- | --- |"]
            for relation in candidate.relations:
                lines.append(f"| {source[relation.left_id]} ↔ {source[relation.right_id]} | {RELATION_NAMES[relation.status]} | "
                             + table_cell("；".join(relation.basis)) + " |")
            if candidate.assignment:
                lines += ["", "## 整体一一对应", "",
                          "每篇新闻固定一个主组，舆论及雷达记录不重复归属。整体比较后选择对应，备选保留在组内。",
                          "对应排序不裁决事实真假；字段冲突仍保留并阻止日报接纳。重新运行可依据新证据调整归属。",
                          "备选对应及排序依据见 [对应候选.csv](对应候选.csv)；完整过程见事件信息.json 的 assignment。"]
        else:
            lines += ["", "## 历史关联评分", "", "该记录来自旧版加权匹配，分数不是正确概率。", ""]
            for key, value in candidate.scores.items():
                lines.append(f"- {key}：{value:.4f}")
        lines += ["", "## 原始事实对照", "", "下表展示各来源提取到的值；“未给出”表示未知，不能补成默认数量或任务。", "",
                  "| 字段 | 新闻 | 舆论 | 雷达 | 核验结论 |", "| --- | --- | --- | --- | --- |"]
        for field in FACT_FIELDS:
            values = [unique([v for o in members if o.source==s for v in o.facts.get(field, [])])
                      for s in ("open_source", "opinion", "radar")]
            verdict = event.verdicts.get(field)
            conclusion = VERDICT_NAMES[verdict.status] if verdict else "尚未核验"
            if verdict and verdict.status == "supported":
                conclusion += "：" + display(verdict.value)
            lines.append("| " + " | ".join(table_cell(v) for v in [FIELD_NAMES[field], *[display(x) for x in values], conclusion]) + " |")
        if event.platform_verdicts:
            lines += ["", "## 联合活动按平台核验", "",
                      "舆论未提及的第二个平台记为未知；各平台数量分别核验，不把单平台数量写成编队总数。", "",
                      "| 平台 | 字段 | 核验结论 |", "| --- | --- | --- |"]
            for platform, verdicts in event.platform_verdicts.items():
                for field, verdict in verdicts.items():
                    lines.append(f"| {table_cell(platform)} | {FIELD_NAMES[field]} | " +
                                 table_cell(VERDICT_NAMES[verdict.status] + ("：" + display(verdict.value)
                                            if verdict.status == "supported" else "")) + " |")
        lines += ["", "## 雷达原始位置", ""]
        for obs in members:
            if obs.source == "radar":
                lines += [f"- `{obs.path}` / {obs.metadata['sheet']} / 批号 `{obs.metadata['batch']}`；"
                          f"原始行号：{obs.metadata['row_numbers']}。"]
                if obs.metadata.get("parent_radar_id"):
                    lines += [f"  - 原批号共有 {len(obs.metadata['batch_row_numbers'])} 条记录，区分为 "
                              f"{obs.metadata['batch_activity_count']} 个活动观测；本组只核验上述行。{obs.metadata['activity_basis']}。"]
                lines += [f"  - 目标身份依据：{obs.metadata.get('identity_basis', '未提供')}。"]
        lines += ["", "## 人工复核", "", "请填写同目录的 `人工复核.csv`。该文件只在首次导出时创建，重复导出会保留填写内容。", ""]
        if stale:
            lines += ["**当前材料或程序核验结果与复核表记录的版本不同。既有意见已保留，请按本次材料重新确认。**", ""]
        lines += ["自动生成的原文副本、雷达子表、关联说明和事件信息会随重新导出更新。", ""]
        return "\n".join(lines)

    def export_day(self, day: DayResult) -> dict:
        has_semantic_review = any(e.semantic_review.get("method") == "llm_semantic_three_source_v1" for e in day.events)
        if has_semantic_review and any(e.classification is None for e in day.events):
            raise ValueError("语义核验快照缺少分类；须重新运行模型核验，不能用规则覆盖")
        if not has_semantic_review and (any(e.classification is None for e in day.events) or not any(
                t.get("agent") == "verification" and t.get("method") == VERIFICATION_METHOD for t in day.trace)):
            # Reverify historical snapshots under the current trust policy.
            rules = next((t.get("rules", {}) for t in day.trace if t.get("agent") == "verification"), {})
            alignment = next((t.get("rules", {}) for t in day.trace if t.get("agent") == "alignment"), {})
            verifier = VerificationAgent(VerificationSettings.model_validate(rules), AlignmentSettings.model_validate(alignment))
            day = day.model_copy(update={"events": verifier.run(day.candidates, day.observations)})
        texts, radar = self.preflight(day)
        original_text_hashes = {hashlib.sha256(content).hexdigest() for content in texts.values()}
        day_dir = self.target(self.output_dir / day.date)
        observations = {o.id: o for o in day.observations}
        events = {e.id: e for e in day.events}
        previous_path = self.target(day_dir / "索引.json")
        previous = json.loads(previous_path.read_text(encoding="utf-8")) if previous_path.exists() else {}
        if previous and (previous.get("format") != FORMAT or previous.get("data_dir") != str(self.data_dir)):
            raise ValueError("已有日期目录包含其他内容或数据集")
        for status_dir in EVENT_STATUSES:
            self.target(day_dir / status_dir).mkdir(parents=True, exist_ok=True)
        main_ids = {c.id for c in day.candidates}
        records = [Candidate(id=stable_id("unpaired_record|" + oid), date=day.date, observation_ids=[oid],
                             relation_status="unmatched", basis=["未找到可靠的跨来源关联；保留原始材料，未配对不代表虚假"])
                   for oid in day.unassigned_observation_ids]
        events.update({e.id: e for e in VerificationAgent(VerificationSettings()).run(records, day.observations)})
        groups = []
        for candidate in day.candidates + records:
            event = events[candidate.id]
            members = [observations[i] for i in candidate.observation_ids]
            status = self.status(candidate, event, members)
            safe_id = candidate.id if re.fullmatch(r"[A-Za-z0-9_-]{1,64}", candidate.id) else stable_id(candidate.id)
            folder_name = ("事件_" if candidate.id in main_ids else "记录_") + safe_id
            folder = self.target(day_dir / status / folder_name)
            old_paths = [self.target(day_dir / s / folder_name) for s in STATUSES if s != status]
            for old in old_paths:
                if old.exists():
                    if folder.exists() or not (old / "事件信息.json").is_file():
                        raise ValueError("分组目录冲突，无法安全保留人工复核记录")
                    identity = json.loads((old / "事件信息.json").read_text(encoding="utf-8"))
                    if identity.get("format") != FORMAT or identity.get("candidate", {}).get("id") != candidate.id:
                        raise ValueError("旧分组目录不属于当前事件，无法安全迁移")
                    folder.parent.mkdir(parents=True, exist_ok=True)
                    old.rename(folder)
            old_meta = {}
            if folder.exists() and any(folder.iterdir()):
                marker = folder / "事件信息.json"
                if not marker.is_file():
                    raise ValueError(f"目录包含非本模块生成的材料：{folder}")
                old_meta = json.loads(marker.read_text(encoding="utf-8"))
                if old_meta.get("format") != FORMAT or old_meta.get("candidate", {}).get("id") != candidate.id:
                    raise ValueError(f"分组目录身份不一致：{folder}")
            folder.mkdir(parents=True, exist_ok=True)
            files = []
            per_source = Counter(o.source for o in members if o.source != "radar")
            ordinal = Counter()
            for obs in members:
                if obs.source == "radar":
                    continue
                ordinal[obs.source] += 1
                filename = SOURCE_NAMES[obs.source] + (f"_{ordinal[obs.source]:02d}" if per_source[obs.source]>1 else "") + ".txt"
                write_bytes(self.target(folder / filename), texts[obs.id])
                files.append({"source": obs.source, "source_name": SOURCE_NAMES[obs.source], "file": filename,
                              "original_path": obs.path, "observation_id": obs.id,
                              "sha256": hashlib.sha256(texts[obs.id]).hexdigest()})
            if any(o.source == "radar" for o in members):
                self.write_radar(self.target(folder / "雷达.xlsx"), members, radar)
            archived = list(old_meta.get("archived_materials", []))
            current_files = {entry["file"] for entry in files}
            if any(o.source == "radar" for o in members):
                current_files.add("雷达.xlsx")
            old_files = {entry["file"] for entry in old_meta.get("copied_texts", [])}
            if any(o.get("source") == "radar" for o in old_meta.get("observations", [])):
                old_files.add("雷达.xlsx")
            for old_file in list(folder.iterdir()):
                if not old_file.is_file() or old_file.name in current_files:
                    continue
                generated_name = bool(re.fullmatch(r"(?:新闻|舆论)(?:_\d+)?\.txt", old_file.name))
                owned = old_file.name in old_files or (generated_name and
                    hashlib.sha256(old_file.read_bytes()).hexdigest() in original_text_hashes)
                if not owned:
                    continue
                source = self.target(old_file)
                revision = stable_id(str(old_meta.get("fingerprint", "legacy")))
                destination = self.target(folder / "历史材料" / f"{revision}_{old_file.name}")
                counter = 1
                while destination.exists():
                    destination = self.target(folder / "历史材料" / f"{revision}_{counter}_{old_file.name}")
                    counter += 1
                destination.parent.mkdir(parents=True, exist_ok=True)
                source.rename(destination)
                archived.append({"file": destination.relative_to(folder).as_posix(),
                                 "former_file": old_file.name, "reason": "not_in_current_assignment"})
            fingerprint = stable_id(json.dumps({"candidate": candidate.model_dump(), "event": event.model_dump(),
                                                "observations": [o.model_dump() for o in members], "files": files},
                                               ensure_ascii=False, sort_keys=True, default=str))
            review = self.target(folder / "人工复核.csv")
            stale = False
            if review.exists():
                stale = review_version(review) != fingerprint
            else:
                write_text(review, csv_text(["事件ID", "数据版本", "关联是否正确", "事实是否可信", "修正后的新闻路径",
                                             "修正后的舆论路径", "修正后的雷达批号", "备注"],
                                            [[candidate.id, fingerprint, "待复核", "待复核", "", "", "", ""]]))
            issues = unique([explain(r) for r in candidate.warnings + event.reasons] +
                            [f"{SOURCE_NAMES[o.source]}：{explain(w)}" for o in members for w in o.warnings
                             if not w.startswith("llm_un")])
            if stale:
                issues.append("现有人工复核表对应旧版本，请重新确认")
            if archived:
                issues.append("旧版未参与当前关联的自动生成材料已保存在历史材料目录")
            write_text(self.target(folder / "关联说明.md"), self.description(candidate, event, members, status, files, issues, stale))
            write_text(self.target(folder / "README.md"), self.group_readme(candidate, event, members, status))
            if candidate.assignment:
                a = candidate.assignment
                rows = [["舆论", v["opinion_id"], v["path"], v["score"], "是" if v["selected"] else "否",
                         a.get("global_margin") if v["selected"] else "", json.dumps(v.get("llm_review"), ensure_ascii=False)]
                        for v in a["opinion_candidates"]]
                rows += [["雷达", v["observation_id"], v.get("batch"), v["score"], "是", v["global_margin"],
                          json.dumps(v["alternatives"], ensure_ascii=False)] for v in a["radar_assignments"]]
                write_text(self.target(folder / "对应候选.csv"), csv_text(
                    ["来源", "证据ID", "路径或批号", "对应排序依据", "当前选择", "整体替代差距", "审查或备选依据"], rows))
            write_json(self.target(folder / "事件信息.json"), {"format": FORMAT, "status": status, "fingerprint": fingerprint,
                       "review_record_stale": stale, "candidate": candidate.model_dump(mode="json"),
                       "event": event.model_dump(mode="json"), "observations": [o.model_dump(mode="json") for o in members],
                       "copied_texts": files, "archived_materials": archived})
            def values(field):
                return display(event.facts.get(field) or unique([v for o in members for v in o.facts.get(field, [])]))
            groups.append({"date": day.date, "status": status, "id": candidate.id,
                           "kind": "main_group" if candidate.id in main_ids else "unassigned_record",
                           "full_sources": {o.source for o in members} == {"open_source", "opinion", "radar"},
                           "accepted": event.accepted, "country": values("country"), "platform": values("platform"),
                           "direction": values("direction"),
                           "news_title": "；".join(o.metadata.get("title", "") for o in members if o.source=="open_source"),
                           "opinion_title": "；".join(o.metadata.get("title", "") for o in members if o.source=="opinion"),
                           "radar_batch": "；".join(o.metadata["batch"] for o in members if o.source=="radar"),
                           "radar_rows": sum(len(o.metadata["row_numbers"]) for o in members if o.source=="radar"),
                           "issues": issues, "path": folder.relative_to(self.output_dir).as_posix(),
                           "relation_status": candidate.relation_status, "basis": candidate.basis, "fact_status": self.fact_status(event),
                           "observation_ids": [o.id for o in members], "source_types": {o.id: o.source for o in members},
                           "assignment": candidate.assignment})
            groups[-1]["original_sources"] = {o.id: {"source": o.source, "path": o.path} for o in members}
        current = {g["path"] for g in groups}
        for old in previous.get("groups", []) + previous.get("unassigned_records", []):
            if old["path"] not in current:
                source = self.target(self.output_dir / old["path"])
                # 已迁移到新状态的目录不存在；失去关联的旧组保留在历史目录。
                if source.exists():
                    if not (source / "事件信息.json").is_file():
                        raise ValueError("历史分组缺少生成标记，无法安全归档")
                    identity = json.loads((source / "事件信息.json").read_text(encoding="utf-8"))
                    if identity.get("format") != FORMAT or identity.get("candidate", {}).get("id") != old["id"]:
                        raise ValueError("失效目录不属于旧分组，无法安全归档")
                    history = self.target(day_dir / "历史分组" / source.name)
                    counter = 1
                    while history.exists():
                        history = self.target(day_dir / "历史分组" / f"{source.name}_{counter}")
                        counter += 1
                    history.parent.mkdir(parents=True, exist_ok=True)
                    source.rename(history)
        # All old category folders are generated paths; remove only empty ones.
        for old_status in OLD_EVENT_STATUSES + ASSOCIATION_STATUSES + LEGACY_STATUSES:
            old_dir = self.target(day_dir / old_status)
            self.remove_old_category_indexes(old_dir)
        unassigned_records = [g for g in groups if g["kind"] == "unassigned_record"]
        groups = [g for g in groups if g["kind"] == "main_group"]
        unassigned = [{"id": oid, "source": observations[oid].source, "path": observations[oid].path,
                       "batch": observations[oid].metadata.get("batch"),
                       "row_numbers": observations[oid].metadata.get("row_numbers", [])}
                      for oid in day.unassigned_observation_ids]
        write_text(self.target(day_dir / "未归属材料.csv"), csv_text(
            ["证据ID", "来源", "原始路径", "雷达批号", "雷达原始行号", "说明"],
            [[v["id"], SOURCE_NAMES[v["source"]], v["path"], v["batch"] or "", str(v["row_numbers"]),
              "未达到对应门槛或未获唯一归属；保留原始材料，不另计主组"] for v in unassigned]))
        index = {"format": FORMAT, "date": day.date, "data_dir": str(self.data_dir), "groups": groups,
                 "unassigned_observations": unassigned,
                 "unassigned_records": unassigned_records,
                 "unpaired_record_count": sum(len(set(g["source_types"].values())) < 2 for g in groups) + len(unassigned_records),
                 "association_rules": next((t.get("rules", {}) for t in day.trace if t.get("agent") == "alignment"), {}),
                 "counts": {s: sum(g["status"] == s for g in groups) for s in EVENT_STATUSES}}
        write_json(previous_path, index)
        for status in EVENT_STATUSES:
            self.write_indexes(day_dir / status, [g for g in groups + unassigned_records if g["status"] == status],
                               [index], category=status)
        self.write_indexes(day_dir, groups, [index])
        logger.info("%s 人工阅读导出：%s", day.date, index["counts"])
        return index

    def remove_old_category_indexes(self, folder: Path):
        """Remove generated navigation after migration; retain manual files."""
        if not folder.is_dir():
            return
        readme = self.target(folder / "README.md")
        if readme.is_file() and readme.read_text(encoding="utf-8").startswith("# 未配对记录\n"):
            readme.unlink()
        index = self.target(folder / "索引.csv")
        if index.is_file():
            with index.open(encoding="utf-8-sig", newline="") as file:
                generated = next(csv.reader(file), []) == INDEX_COLUMNS
            if generated:
                index.unlink()
        if not any(folder.iterdir()):
            folder.rmdir()

    def write_indexes(self, folder: Path, groups: list[dict], days: list[dict], category: str | None = None):
        rows = [[g["date"], g["status"], g["id"], "是" if g["full_sources"] else "否", "是" if g["accepted"] else "否",
                 g["country"], g["platform"], g["direction"], g["news_title"], g["opinion_title"], g["radar_batch"],
                 g["radar_rows"], "；".join(g["issues"]), g["path"], "；".join(g.get("basis", [])),
                 g.get("fact_status", "历史记录"), g.get("assignment", {}).get("news_number", ""),
                 g.get("assignment", {}).get("opinion_number") or "",
                 "每篇新闻一个主组" if g.get("assignment") else "单来源记录" if g.get("kind") == "unassigned_record" else "历史事件组"] for g in groups]
        write_text(self.target(folder / "索引.csv"), csv_text(INDEX_COLUMNS, rows))
        counts = Counter(g["status"] for g in groups)
        one_to_one = bool(groups) and all(g.get("assignment") for g in groups)
        statuses = EVENT_STATUSES
        extra_records = sum(len(d.get("unassigned_records", [])) for d in days)
        if category:
            main_groups = sum(g.get("kind") != "unassigned_record" for g in groups)
            lines = [f"# {category}", "", f"共 {main_groups} 个主组、{len(groups) - main_groups} 条未归属记录。", "",
                     "可信要求三类来源完成同一事件关联、关键事实至少两类来源支持且第三类无明确冲突。"
                     "冲突、缺源、关联歧义、证据不足及未配对记录均归入不可信，并保留原始材料和具体原因。", ""]
        else:
            lines = ["# 事件核验分组", "", f"共 {len(days)} 天、{len(groups)} 个{'新闻主组' if one_to_one else '事件主组'}；日报接纳 {sum(g['accepted'] for g in groups)} 项。", "",
                     "| 分类 | 主组数 | 判断标准 |", "| --- | --- | --- |",
                     f"| 可信事件 | {counts['可信事件']} | 三类来源完成同一事件关联，关键事实至少两类来源支持，第三类无明确冲突值 |",
                     f"| 不可信事件 | {counts['不可信事件']} | 除可信事件外的全部主组，包括冲突、缺源、关联歧义、证据不足和未配对 |", "",
                     f"另有 {extra_records} 条未归属的舆论或雷达，存入“不可信事件”，不增加新闻主组数量。", "",
                     "## 分组依据", "",
                     "- **先关联，再核验**：新闻、舆论、雷达三类来源须完成同一事件关联；其余均归入不可信事件，并注明具体原因。",
                     "- **来源缺失**：缺失记未知；即使已有两类来源一致，缺少第三类关联材料仍归入不可信。第三类已关联但未给出某字段时，不构成该字段冲突。",
                     "- **逐字段核验**：可信事件仅保留已交叉印证的字段；数量或任务只有单类来源支持时记未知，不补全。",
                     "- **雷达依据**：雷达记录作为时空、数量和轨迹核验依据之一，不作为无条件优先真值；任一来源的明确冲突均不能被另外两类来源一致抵消。",
                     "- **目标与活动范围**：飞机、舰艇等均先区分具体目标和活动；批号、型号不自动代表同一目标。批号内不同活动分别关联，明确目标的异常移动及同一活动内的事实冲突仍保留。",
                     "- **时空与歧义**：同一天、同方向不足以证明同一活动；数小时差距或多个合理对应仅保留候选。",
                     "- **日报范围**：可信事件与日报接纳分别记录，核心海域以外的可信活动也保留在可信事件中。", ""]
            if one_to_one:
                lines += ["每篇新闻固定一个主组，舆论及雷达最多归属一个主组；新闻、舆论、雷达缺失时不强行凑齐。"
                          "联合活动按平台分别核验，仍计一个主组；备选对应保留在组内，不另计事件。", ""]
        lines += [
                  "先打开 `索引.csv` 筛选，再读每组 `README.md` 的依据；完整字段见 `关联说明.md`。" +
                  ("原文按字节复制，主组来源不重复；雷达子表包含该主组对应活动的原始行，同批号其他活动分别关联或保留为未归属材料。" if one_to_one else
                   "原文按字节复制，雷达只筛对应记录；同一原文可因子事件或多个候选出现在不同组中。"), "",
                  "人工意见填写 `人工复核.csv`，重复导出保留；意见暂不自动重算。", ""]
        if not category:
            lines += ["## 日期导航", "", "| 日期 | " + " | ".join(statuses) + " | 不可信未归属记录 |",
                      "| --- | " + " | ".join("---" for _ in statuses) + " | --- |"]
            for d in days:
                target = self.output_dir / d["date"] / "README.md"
                relative = Path(relpath(target, folder)).as_posix()
                lines.append(f"| [{d['date']}](<{quote(relative)}>) | "
                             + " | ".join(str(d['counts'].get(s, 0)) for s in statuses)
                             + f" | {len(d.get('unassigned_records', []))} |")
        if len(days) == 1 and folder != self.output_dir:
            lines += ["", "## 事件导航", "", "| 分类 | 国家地区 | 平台 | 方向 | 材料 |", "| --- | --- | --- | --- | --- |"]
            for g in groups:
                relative = Path(relpath(self.output_dir / g["path"] / "README.md", folder)).as_posix()
                lines.append("| " + " | ".join(table_cell(g[k]) for k in ("status", "country", "platform", "direction")) +
                             f" | [打开分组](<{quote(relative)}>) |")
        write_text(self.target(folder / "README.md"), "\n".join(lines) + "\n")

    def export(self, days: Iterable[DayResult]) -> dict:
        updated = []
        for day in days:
            self.export_day(day)
            updated.append(day.date)
        if not updated:
            raise ValueError("没有可导出的日关联结果，请先运行 run")
        indexes = []
        for path in sorted(self.output_dir.glob("????-??-??/索引.json")):
            path = self.target(path)
            info = json.loads(path.read_text(encoding="utf-8"))
            if info.get("format") != FORMAT or info.get("data_dir") != str(self.data_dir):
                raise ValueError("已有日期索引不属于当前数据集")
            indexes.append(info)
        groups = [g for d in indexes for g in d["groups"]]
        records = [g for d in indexes for g in d.get("unassigned_records", [])]
        observations = {oid: source for g in groups for oid, source in g["source_types"].items()}
        observations.update({v["id"]: v["source"] for d in indexes for v in d.get("unassigned_observations", [])})
        self.write_indexes(self.output_dir, groups, indexes)
        for status in EVENT_STATUSES:
            self.write_indexes(self.output_dir / status, [g for g in groups + records if g["status"] == status],
                               indexes, category=status)
        for old_status in OLD_EVENT_STATUSES:
            self.remove_old_category_indexes(self.target(self.output_dir / old_status))
        manifest = {"format": FORMAT, "data_dir": str(self.data_dir), "output_dir": str(self.output_dir),
                    "updated_dates": sorted(set(updated)), "exported_dates": [d["date"] for d in indexes],
                    "total_groups": len(groups), "counts": {s: sum(g["status"]==s for g in groups)
                        for s in EVENT_STATUSES},
                    "unique_observations": dict(Counter(observations.values()))}
        manifest["unassigned_observations"] = sum(len(d.get("unassigned_observations", [])) for d in indexes)
        manifest["unpaired_record_count"] = sum(len(set(g["source_types"].values())) < 2 for g in groups) + len(records)
        manifest["unassigned_record_sources"] = dict(Counter(s for g in records for s in g["source_types"].values()))
        originals = {(v["source"], v["path"]) for g in groups for v in g.get("original_sources", {}).values()}
        originals.update((v["source"], v["path"]) for d in indexes for v in d.get("unassigned_observations", []))
        manifest["unique_original_files"] = dict(Counter(s for s, _ in originals))
        write_json(self.target(self.output_dir / "manifest.json"), manifest)
        return manifest


def export_saved_results(data_dir: Path, analysis_dir: Path, output_dir: Path, dates: list[str] | None = None) -> dict:
    if not analysis_dir.is_dir():
        raise FileNotFoundError(f"没有已保存的日关联结果：{analysis_dir}；请先运行 run")
    paths = sorted(analysis_dir.glob("????-??-??.json")) if dates is None else [analysis_dir / (d + ".json") for d in sorted(set(dates))]
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(f"关联结果缺失：{path}；请先运行对应日期的 run")
    exporter = ReviewExporter(data_dir, output_dir, analysis_dir)
    def load_days():
        for path in paths:
            day = DayResult.model_validate_json(path.read_text(encoding="utf-8"))
            if day.date != path.stem:
                raise ValueError(f"关联结果日期与文件名不一致：{path}")
            yield day
    return exporter.export(load_days())
