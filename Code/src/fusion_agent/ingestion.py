from collections import defaultdict
import hashlib
from pathlib import Path
import re

import openpyxl

from .models import Document, RadarBatch
from .utils import parse_date, stable_id


class IngestionAgent:
    name = "ingestion"

    def __init__(self, root: Path):
        self.root = root

    def dates(self) -> dict[str, Path]:
        if not self.root.is_dir():
            raise FileNotFoundError(f"数据目录不存在：{self.root}")
        result = {}
        for path in self.root.iterdir():
            if path.is_dir() and re.fullmatch(r"\d{4}\.\d{1,2}\.\d{1,2}", path.name):
                day = parse_date(path.name)
                if day in result:
                    raise ValueError(f"重复日期目录：{day}")
                result[day] = path
        if not result:
            raise ValueError(f"未找到按日期组织的数据：{self.root}")
        return dict(sorted(result.items()))

    def fingerprint(self, day: str) -> str:
        h = hashlib.sha256()
        for path in sorted(self.dates()[day].rglob("*")):
            if path.is_file() and not path.name.startswith("~$") and path.suffix.lower() in {".txt", ".xlsx"}:
                h.update(str(path.relative_to(self.root)).encode("utf-8"))
                h.update(path.read_bytes())
        return h.hexdigest()

    def run(self, day: str) -> tuple[list[Document], list[RadarBatch], dict]:
        folder = self.dates()[day]
        docs = []
        counts = {}
        for source, name in (("open_source", "开源军情"), ("opinion", "网上舆论")):
            paths = sorted((folder / name).glob("*.txt"))
            if not paths:
                raise ValueError(f"{day} 缺少 {name} 文本")
            counts[source] = len(paths)
            for path in paths:
                text = path.read_text(encoding="utf-8-sig")
                title = re.search(r"^标题[：:]\s*(.*)$", text, re.MULTILINE)
                relative = path.relative_to(self.root).as_posix()
                docs.append(Document(id=stable_id(relative), date=day, source=source,
                                     path=relative, text=text, title=title.group(1) if title else ""))
        workbooks = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]
        if len(workbooks) != 1:
            raise ValueError(f"{day} 应有一个 XLSX，实际为 {len(workbooks)}")
        path = workbooks[0]
        relative = path.relative_to(self.root).as_posix()
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        batches = []
        try:
            for sheet_name in ("飞机", "舰船"):
                if sheet_name not in wb.sheetnames:
                    raise ValueError(f"{path.name} 缺少工作表：{sheet_name}")
                rows = wb[sheet_name].iter_rows(values_only=True)
                headers = next(rows)
                required = ({"批号", "国家地区", "目标类型", "位置时间", "经度", "纬度", "目标数量", "任务"}
                            if sheet_name == "飞机" else
                            {"批号", "国家地区", "机舰名", "活动区域", "首点时间", "结束时间", "数量", "任务"})
                if not required <= set(headers):
                    raise ValueError(f"{path.name}/{sheet_name} 缺少字段：{required-set(headers)}")
                grouped = defaultdict(list)
                for number, values in enumerate(rows, 2):
                    if all(v is None for v in values):
                        continue
                    row = dict(zip(headers, values))
                    if not row.get("批号"):
                        raise ValueError(f"{path.name}/{sheet_name} 第 {number} 行缺少批号")
                    grouped[str(row["批号"])].append((number, row))
                counts[sheet_name + "_rows"] = sum(len(rs) for rs in grouped.values())
                for batch, rs in sorted(grouped.items()):
                    batches.append(RadarBatch(id=stable_id(relative + sheet_name + batch), date=day,
                                              path=relative, sheet=sheet_name, batch=batch,
                                              rows=[r for _, r in rs], row_numbers=[n for n, _ in rs]))
        finally:
            wb.close()
        counts["radar_batches"] = len(batches)
        return docs, batches, counts

